"""Legacy-монолит приложения.

Основная поддерживаемая точка входа — ``main.py`` с вкладками из каталога
``modules/``. Этот файл сохранен для справки и обратной совместимости; новую
логику следует добавлять в модульную версию, чтобы не поддерживать две копии
одного и того же кода.
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import os
import sys
import threading
import re
import nltk
from tqdm import tqdm
from bs4 import BeautifulSoup
import numpy as np
from collections import Counter
import matplotlib.pyplot as plt
from matplotlib.figure import Figure
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from wordcloud import WordCloud

# Загружаем необходимые ресурсы NLTK
try:
    nltk.data.find('tokenizers/punkt')
except LookupError:
    nltk.download('punkt', quiet=True)

class MultiTextApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Многофункциональное приложение для обработки текста")
        self.geometry("900x700")  # Увеличиваем размер окна по умолчанию для лучшего отображения
        self.resizable(True, True)
        
        # Устанавливаем минимальный размер окна для лучшей адаптивности
        self.minsize(800, 600)
        
        # Создаем систему вкладок
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill='both', expand=True, padx=10, pady=10)
        
        # Добавляем все вкладки с функциональностью
        self.fb2_converter_tab = FB2ConverterTab(self.notebook)
        self.notebook.add(self.fb2_converter_tab, text="Конвертация в FB2")
        
        self.reading_time_tab = ReadingTimeTab(self.notebook)
        self.notebook.add(self.reading_time_tab, text="Время чтения")
        
        self.docx_converter_tab = DocxConverterTab(self.notebook)
        self.notebook.add(self.docx_converter_tab, text="Конвертация DOCX")
        
        self.text_analyzer_tab = TextAnalyzerTab(self.notebook)
        self.notebook.add(self.text_analyzer_tab, text="Анализ текста")
        
        self.text_splitter_tab = TextSplitterTab(self.notebook)
        self.notebook.add(self.text_splitter_tab, text="Разделение текста")
        
        # Вкладка статистического анализа текста
        self.text_statistics_tab = TextStatisticsTab(self.notebook)
        self.notebook.add(self.text_statistics_tab, text="Статистика текста")
        
        # Создаем строку состояния
        self.status_var = tk.StringVar()
        self.status_var.set("Готово")
        self.status_bar = ttk.Label(self, textvariable=self.status_var, relief=tk.SUNKEN, anchor=tk.W)
        self.status_bar.pack(side=tk.BOTTOM, fill=tk.X)
        
        # Настройка стилей
        self.style = ttk.Style()
        self.style.configure("TButton", padding=6, relief="flat", font=("Helvetica", 10))
        self.style.configure("TLabel", font=("Helvetica", 10))
        self.style.configure("TEntry", font=("Helvetica", 10))
        
        # Установка обработчика закрытия окна
        self.protocol("WM_DELETE_WINDOW", self.on_closing)
    
    def update_status(self, message):
        """Обновляет сообщение в строке состояния"""
        self.status_var.set(message)
        self.update_idletasks()
    
    def on_closing(self):
        """Обработчик закрытия окна"""
        if messagebox.askokcancel("Выход", "Вы действительно хотите выйти?"):
            self.destroy()


class FB2ConverterTab(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        
        # Создаем элементы интерфейса
        self.create_widgets()
    
    def create_widgets(self):
        # Фрейм для выбора файла
        file_frame = ttk.LabelFrame(self, text="Выбор файла")
        file_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        # Поле для отображения выбранного файла
        self.file_var = tk.StringVar()
        self.file_entry = ttk.Entry(file_frame, textvariable=self.file_var, width=60)
        self.file_entry.grid(row=0, column=0, padx=5, pady=5, sticky="ew")
        
        # Кнопка для выбора файла
        self.browse_button = ttk.Button(file_frame, text="Обзор", command=self.browse_file)
        self.browse_button.grid(row=0, column=1, padx=5, pady=5)
        
        # Фрейм для конвертации файлов
        options_frame = ttk.LabelFrame(self, text="Параметры конвертации")
        options_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        # Радиокнопки для выбора типа конвертации
        self.convert_type = tk.StringVar(value="txt_to_fb2")
        
        self.txt_to_fb2_radio = ttk.Radiobutton(
            options_frame, text="TXT в FB2", 
            variable=self.convert_type, value="txt_to_fb2"
        )
        self.txt_to_fb2_radio.grid(row=0, column=0, padx=5, pady=5, sticky="w")
        
        self.docx_to_fb2_radio = ttk.Radiobutton(
            options_frame, text="DOCX в FB2", 
            variable=self.convert_type, value="docx_to_fb2"
        )
        self.docx_to_fb2_radio.grid(row=1, column=0, padx=5, pady=5, sticky="w")
        
        # Кнопка для запуска конвертации
        self.convert_button = ttk.Button(
            self, text="Конвертировать", 
            command=self.start_conversion
        )
        self.convert_button.pack(pady=10)
        
        # Индикатор прогресса
        self.progress = ttk.Progressbar(self, orient="horizontal", length=300, mode="indeterminate")
        self.progress.pack(pady=5, fill="x", padx=10)
        
        # Текстовое поле для вывода статуса операции
        self.status_text = tk.Text(self, height=10, width=70, wrap="word")
        self.status_text.pack(padx=10, pady=10, fill="both", expand=True)
        
        # Настройка scrollbar для текстового поля
        scrollbar = ttk.Scrollbar(self.status_text, command=self.status_text.yview)
        scrollbar.pack(side="right", fill="y")
        self.status_text.config(yscrollcommand=scrollbar.set)
        
        # Делаем первую колонку растягиваемой
        file_frame.columnconfigure(0, weight=1)
    
    def browse_file(self):
        """Открывает диалог выбора файла"""
        convert_type = self.convert_type.get()
        
        if convert_type == "txt_to_fb2":
            filetypes = [("Текстовые файлы", "*.txt")]
        elif convert_type == "docx_to_fb2":
            filetypes = [("DOCX файлы", "*.docx")]
        
        filename = filedialog.askopenfilename(filetypes=filetypes)
        if filename:
            self.file_var.set(filename)
    
    def start_conversion(self):
        """Запускает процесс конвертации в отдельном потоке"""
        file_path = self.file_var.get()
        if not file_path:
            messagebox.showerror("Ошибка", "Выберите файл для конвертации!")
            return
        
        # Проверяем существование файла
        if not os.path.exists(file_path):
            messagebox.showerror("Ошибка", "Указанный файл не существует!")
            return
        
        # Очищаем поле статуса
        self.status_text.delete(1.0, tk.END)
        self.update_status("Начало конвертации...")
        
        # Запускаем прогресс-бар
        self.progress.start()
        
        # Запускаем конвертацию в отдельном потоке
        convert_thread = threading.Thread(
            target=self.perform_conversion,
            args=(file_path,)
        )
        convert_thread.daemon = True
        convert_thread.start()
    
    def perform_conversion(self, file_path):
        """Выполняет конвертацию файла"""
        try:
            convert_type = self.convert_type.get()
            
            if convert_type == "txt_to_fb2":
                output_file = self.txt_to_fb2(file_path)
            elif convert_type == "docx_to_fb2":
                output_file = self.docx_to_fb2(file_path)
            
            # Отображаем успешное завершение
            self.after(100, lambda: self.conversion_completed(output_file))
            
        except Exception as e:
            # Обрабатываем ошибки
            error_message = f"Ошибка конвертации: {str(e)}"
            self.after(100, lambda: self.update_status(error_message))
            self.after(100, lambda: self.progress.stop())
    
    def conversion_completed(self, output_file):
        """Обрабатывает завершение конвертации"""
        self.progress.stop()
        self.update_status(f"Конвертация завершена успешно!\nРезультат сохранен в: {output_file}")
        
        if messagebox.askyesno("Конвертация завершена", 
                               "Конвертация завершена успешно. Открыть файл?"):
            # Открытие файла в браузере или программе по умолчанию
            output_file_url = f'file:///{os.path.normpath(os.path.abspath(output_file)).replace(os.sep, "/")}'
            os.system(f'start "" "{output_file_url}"')
    
    def update_status(self, message):
        """Обновляет текстовое поле статуса"""
        self.status_text.insert(tk.END, message + "\n")
        self.status_text.see(tk.END)  # Прокрутка вниз
    
    def txt_to_fb2(self, txt_file_path):
        """Конвертирует TXT файл в FB2 формат"""
        self.update_status("Чтение текстового файла...")
        
        # Чтение текстового файла
        with open(txt_file_path, "r", encoding="utf-8") as f:
            text = f.read()
        
        # Токенизация текста
        self.update_status("Обработка текста...")
        tokens = nltk.word_tokenize(text)
        
        lines = []
        current_line = []
        for token in tokens:
            current_line.append(token)
            if len(current_line) >= 25:
                lines.append(' '.join(current_line))
                current_line = []
        
        if current_line:
            lines.append(' '.join(current_line))
        
        # Получаем имя файла без расширения
        file_name = os.path.splitext(os.path.basename(txt_file_path))[0]
        
        # Формируем имя выходного файла
        output_file_name = f"{file_name}_formatted.fb2"
        
        # Генерация FB2 файла через BeautifulSoup
        self.update_status("Создание FB2 файла...")
        
        soup = BeautifulSoup(features='xml')
        soup.append(soup.new_tag("FictionBook", xmlns="http://www.gribuser.ru/xml/fictionbook/2.0"))
        fiction_book = soup.FictionBook
        
        # Метаданные
        description = soup.new_tag("description")
        title_info = soup.new_tag("title-info")
        
        book_title = soup.new_tag("book-title")
        book_title.string = file_name
        title_info.append(book_title)
        
        author = soup.new_tag("author")
        author_name = soup.new_tag("first-name")
        author_name.string = "Автор"
        author.append(author_name)
        title_info.append(author)
        
        description.append(title_info)
        fiction_book.append(description)
        
        # Тело документа
        body = soup.new_tag("body")
        section = soup.new_tag("section")
        
        for line in lines:
            p = soup.new_tag("p")
            p.string = line
            section.append(p)
        
        body.append(section)
        fiction_book.append(body)
        
        # Сохраняем результат в FB2
        with open(output_file_name, 'w', encoding='utf-8') as file:
            file.write(str(soup))
        
        self.update_status(f"FB2 файл сохранен как '{output_file_name}'")
        return output_file_name
    
    def docx_to_fb2(self, docx_file_path):
        """Конвертирует DOCX файл в FB2 формат"""
        self.update_status("Чтение DOCX файла...")
        
        try:
            from docx import Document
            from xml.sax.saxutils import escape
            
            # Открываем DOCX файл
            doc = Document(docx_file_path)
            
            # Получаем имя файла без расширения
            file_name = os.path.splitext(os.path.basename(docx_file_path))[0]
            
            # Формируем имя выходного файла
            output_file_name = f"{file_name}.fb2"
            
            self.update_status("Создание FB2 файла...")
            
            # Открываем файл для записи в формате FB2
            with open(output_file_name, "w", encoding="utf-8") as fb2_file:
                # Записываем начало документа FB2
                fb2_file.write('<?xml version="1.0" encoding="UTF-8"?>\n')
                fb2_file.write('<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0" xmlns:l="http://www.w3.org/1999/xlink">\n')
                
                # Записываем метаинформацию о книге
                fb2_file.write('<description>\n')
                fb2_file.write('<title-info>\n')
                fb2_file.write(f'<book-title>{escape(file_name)}</book-title>\n')
                fb2_file.write('<author>\n')
                fb2_file.write('<first-name>Автор</first-name>\n')
                fb2_file.write('<last-name></last-name>\n')
                fb2_file.write('</author>\n')
                fb2_file.write('</title-info>\n')
                fb2_file.write('</description>\n')
                
                fb2_file.write('<body>\n')
                
                # Перебираем все параграфы в документе DOCX
                for paragraph in doc.paragraphs:
                    # Экранируем специальные символы в тексте и записываем его в FB2
                    fb2_file.write('<p>' + escape(paragraph.text) + '</p>\n')
                
                # Закрываем документ FB2
                fb2_file.write('</body>\n')
                fb2_file.write('</FictionBook>')
            
            self.update_status(f"FB2 файл сохранен как '{output_file_name}'")
            return output_file_name
            
        except ImportError:
            self.update_status("Ошибка: модуль python-docx не установлен")
            raise Exception("Для конвертации DOCX файлов требуется модуль python-docx. Установите его командой: pip install python-docx")


class ReadingTimeTab(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        
        # Создаем элементы интерфейса
        self.create_widgets()
    
    def create_widgets(self):
        # Фрейм для выбора файла
        file_frame = ttk.LabelFrame(self, text="Выбор файла")
        file_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        # Поле для отображения выбранного файла
        self.file_var = tk.StringVar()
        self.file_entry = ttk.Entry(file_frame, textvariable=self.file_var, width=60)
        self.file_entry.grid(row=0, column=0, padx=5, pady=5, sticky="ew")
        
        # Кнопка для выбора файла
        self.browse_button = ttk.Button(file_frame, text="Обзор", command=self.browse_file)
        self.browse_button.grid(row=0, column=1, padx=5, pady=5)
        
        # Фрейм для настроек расчета времени чтения
        settings_frame = ttk.LabelFrame(self, text="Настройки")
        settings_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        # Настройка скорости чтения (слов в минуту)
        ttk.Label(settings_frame, text="Скорость чтения (слов в минуту):").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        
        self.wpm_var = tk.IntVar(value=200)
        self.wpm_spinbox = ttk.Spinbox(settings_frame, from_=50, to=1000, increment=10, textvariable=self.wpm_var, width=5)
        self.wpm_spinbox.grid(row=0, column=1, padx=5, pady=5, sticky="w")
        
        # Кнопка для расчета времени чтения
        self.calculate_button = ttk.Button(
            self, text="Рассчитать время чтения", 
            command=self.start_calculation
        )
        self.calculate_button.pack(pady=10)
        
        # Индикатор прогресса
        self.progress = ttk.Progressbar(self, orient="horizontal", length=300, mode="indeterminate")
        self.progress.pack(pady=5, fill="x", padx=10)
        
        # Рамка для отображения результатов
        results_frame = ttk.LabelFrame(self, text="Результаты")
        results_frame.pack(fill="both", padx=10, pady=10, expand=True)
        
        # Поля для отображения результатов
        self.total_words_var = tk.StringVar(value="Всего слов: 0")
        self.reading_time_var = tk.StringVar(value="Время чтения: 0 мин. 0 сек.")
        
        ttk.Label(results_frame, textvariable=self.total_words_var, font=("Helvetica", 12)).pack(pady=10)
        ttk.Label(results_frame, textvariable=self.reading_time_var, font=("Helvetica", 12)).pack(pady=10)
        
        # Делаем первую колонку растягиваемой
        file_frame.columnconfigure(0, weight=1)
        settings_frame.columnconfigure(0, weight=1)
    
    def browse_file(self):
        """Открывает диалог выбора файла"""
        filetypes = [
            ("Текстовые файлы", "*.txt"), 
            ("Все файлы", "*.*")
        ]
        
        filename = filedialog.askopenfilename(filetypes=filetypes)
        if filename:
            self.file_var.set(filename)
    
    def start_calculation(self):
        """Запускает расчет времени чтения в отдельном потоке"""
        file_path = self.file_var.get()
        if not file_path:
            messagebox.showerror("Ошибка", "Выберите файл для анализа!")
            return
        
        # Проверяем существование файла
        if not os.path.exists(file_path):
            messagebox.showerror("Ошибка", "Указанный файл не существует!")
            return
        
        # Очищаем результаты
        self.total_words_var.set("Всего слов: расчет...")
        self.reading_time_var.set("Время чтения: расчет...")
        
        # Получаем скорость чтения
        try:
            wpm = int(self.wpm_var.get())
            if wpm <= 0:
                raise ValueError("Скорость чтения должна быть положительным числом")
        except ValueError:
            messagebox.showerror("Ошибка", "Пожалуйста, введите корректное значение для скорости чтения!")
            return
        
        # Запускаем прогресс-бар
        self.progress.start()
        
        # Запускаем расчет в отдельном потоке
        calc_thread = threading.Thread(
            target=self.calculate_reading_time,
            args=(file_path, wpm)
        )
        calc_thread.daemon = True
        calc_thread.start()
    
    def calculate_reading_time(self, file_path, words_per_minute):
        """Выполняет расчет времени чтения"""
        try:
            # Читаем файл
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
            
            # Разбиваем текст на слова
            words = text.split()
            
            # Подсчитываем количество слов в тексте
            num_words = len(words)
            
            # Рассчитываем время чтения в минутах
            reading_time_minutes = num_words / words_per_minute
            
            # Преобразуем время в формат минут и секунд
            minutes = int(reading_time_minutes)
            seconds = int((reading_time_minutes - minutes) * 60)
            
            # Обновляем интерфейс в основном потоке
            self.after(100, lambda: self.update_results(num_words, minutes, seconds))
            
        except Exception as e:
            # Обрабатываем ошибки
            error_message = f"Ошибка при расчете: {str(e)}"
            self.after(100, lambda: messagebox.showerror("Ошибка", error_message))
            self.after(100, lambda: self.progress.stop())
    
    def update_results(self, num_words, minutes, seconds):
        """Обновляет отображение результатов"""
        self.progress.stop()
        self.total_words_var.set(f"Всего слов: {num_words}")
        self.reading_time_var.set(f"Время чтения: {minutes} мин. {seconds} сек.")


class DocxConverterTab(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        
        # Создаем элементы интерфейса
        self.create_widgets()
    
    def create_widgets(self):
        # Фрейм для выбора файла
        file_frame = ttk.LabelFrame(self, text="Выбор файла")
        file_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        # Поле для отображения выбранного файла
        self.file_var = tk.StringVar()
        self.file_entry = ttk.Entry(file_frame, textvariable=self.file_var, width=60)
        self.file_entry.grid(row=0, column=0, padx=5, pady=5, sticky="ew")
        
        # Кнопка для выбора файла
        self.browse_button = ttk.Button(file_frame, text="Обзор", command=self.browse_file)
        self.browse_button.grid(row=0, column=1, padx=5, pady=5)
        
        # Кнопка для конвертации
        self.convert_button = ttk.Button(
            self, text="Конвертировать DOCX в TXT", 
            command=self.start_conversion
        )
        self.convert_button.pack(pady=10)
        
        # Индикатор прогресса
        self.progress = ttk.Progressbar(self, orient="horizontal", length=300, mode="indeterminate")
        self.progress.pack(pady=5, fill="x", padx=10)
        
        # Текстовое поле для вывода статуса операции
        self.status_text = tk.Text(self, height=10, width=70, wrap="word")
        self.status_text.pack(padx=10, pady=10, fill="both", expand=True)
        
        # Настройка scrollbar для текстового поля
        scrollbar = ttk.Scrollbar(self.status_text, command=self.status_text.yview)
        scrollbar.pack(side="right", fill="y")
        self.status_text.config(yscrollcommand=scrollbar.set)
        
        # Делаем первую колонку растягиваемой
        file_frame.columnconfigure(0, weight=1)
    
    def browse_file(self):
        """Открывает диалог выбора файла"""
        filetypes = [("Word Files", "*.docx")]
        
        filename = filedialog.askopenfilename(filetypes=filetypes)
        if filename:
            self.file_var.set(filename)
    
    def start_conversion(self):
        """Запускает процесс конвертации в отдельном потоке"""
        file_path = self.file_var.get()
        if not file_path:
            messagebox.showerror("Ошибка", "Выберите файл для конвертации!")
            return
        
        # Проверяем существование файла
        if not os.path.exists(file_path):
            messagebox.showerror("Ошибка", "Указанный файл не существует!")
            return
        
        # Проверяем, является ли файл действительно DOCX
        if not file_path.lower().endswith('.docx'):
            messagebox.showerror("Ошибка", "Выбранный файл не является файлом DOCX!")
            return
        
        # Очищаем поле статуса
        self.status_text.delete(1.0, tk.END)
        self.update_status("Начало конвертации...")
        
        # Запускаем прогресс-бар
        self.progress.start()
        
        # Запускаем конвертацию в отдельном потоке
        convert_thread = threading.Thread(
            target=self.perform_conversion,
            args=(file_path,)
        )
        convert_thread.daemon = True
        convert_thread.start()
    
    def perform_conversion(self, file_path):
        """Выполняет конвертацию файла"""
        try:
            # Пытаемся импортировать модуль docx2txt
            try:
                import docx2txt
            except ImportError:
                self.after(100, lambda: self.update_status("Ошибка: модуль docx2txt не установлен"))
                self.after(100, lambda: self.progress.stop())
                raise Exception("Для конвертации DOCX файлов требуется модуль docx2txt. Установите его командой: pip install docx2txt")
            
            # Конвертируем DOCX в текст
            self.update_status("Конвертация DOCX в текст...")
            text = docx2txt.process(file_path)
            
            # Получаем имя файла без расширения
            file_name = os.path.splitext(os.path.basename(file_path))[0]
            
            # Формируем имя выходного файла
            output_file_name = f"{file_name}.txt"
            
            # Сохраняем результат в TXT
            with open(output_file_name, 'w', encoding='utf-8') as file:
                file.write(text)
            
            # Отображаем успешное завершение
            self.after(100, lambda: self.conversion_completed(output_file_name))
            
        except Exception as e:
            # Обрабатываем ошибки
            error_message = f"Ошибка конвертации: {str(e)}"
            self.after(100, lambda: self.update_status(error_message))
            self.after(100, lambda: self.progress.stop())
    
    def conversion_completed(self, output_file):
        """Обрабатывает завершение конвертации"""
        self.progress.stop()
        self.update_status(f"Конвертация завершена успешно!\nРезультат сохранен в: {output_file}")
        
        if messagebox.askyesno("Конвертация завершена", 
                               "Конвертация завершена успешно. Открыть файл?"):
            # Открываем файл в блокноте или текстовом редакторе по умолчанию
            os.system(f'notepad.exe "{output_file}"')
    
    def update_status(self, message):
        """Обновляет текстовое поле статуса"""
        self.status_text.insert(tk.END, message + "\n")
        self.status_text.see(tk.END)  # Прокрутка вниз


class WordSearchFrame(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        
        # Создаем элементы интерфейса поиска слов
        self.create_widgets()
    
    def create_widgets(self):
        # Фрейм для выбора файла
        file_frame = ttk.LabelFrame(self, text="Выбор файла")
        file_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        # Поле для отображения выбранного файла
        self.file_var = tk.StringVar()
        self.file_entry = ttk.Entry(file_frame, textvariable=self.file_var, width=60)
        self.file_entry.grid(row=0, column=0, padx=5, pady=5, sticky="ew")
        
        # Кнопка для выбора файла
        self.browse_button = ttk.Button(file_frame, text="Обзор", command=self.browse_file)
        self.browse_button.grid(row=0, column=1, padx=5, pady=5)
        
        # Фрейм для ввода слова поиска
        search_frame = ttk.LabelFrame(self, text="Поиск")
        search_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        ttk.Label(search_frame, text="Введите слово или его часть для поиска:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        
        self.search_var = tk.StringVar()
        self.search_entry = ttk.Entry(search_frame, textvariable=self.search_var, width=40)
        self.search_entry.grid(row=1, column=0, padx=5, pady=5, sticky="ew")
        
        # Кнопка для запуска поиска
        self.search_button = ttk.Button(
            search_frame, text="Найти", 
            command=self.start_search
        )
        self.search_button.grid(row=1, column=1, padx=5, pady=5)
        
        # Кнопка сброса
        self.clear_button = ttk.Button(
            search_frame, text="Сброс", 
            command=self.clear_search
        )
        self.clear_button.grid(row=1, column=2, padx=5, pady=5)
        
        # Результаты поиска
        result_frame = ttk.LabelFrame(self, text="Результаты поиска")
        result_frame.pack(fill="both", padx=10, pady=10, expand=True)
        
        self.result_text = tk.Text(result_frame, wrap="word", height=10)
        self.result_text.pack(fill="both", expand=True, padx=5, pady=5)
        
        # Scrollbar для текста результатов
        scrollbar = ttk.Scrollbar(self.result_text, command=self.result_text.yview)
        scrollbar.pack(side="right", fill="y")
        self.result_text.config(yscrollcommand=scrollbar.set)
        
        # Фрейм для кнопок управления результатами
        buttons_frame = ttk.Frame(result_frame)
        buttons_frame.pack(fill="x", padx=5, pady=5)
        
        # Кнопка копирования результатов
        self.copy_button = ttk.Button(
            buttons_frame, text="Копировать результаты", 
            command=self.copy_results
        )
        self.copy_button.pack(side=tk.LEFT, padx=5)
        
        # Кнопка сохранения результатов в файл
        self.save_button = ttk.Button(
            buttons_frame, text="Сохранить в файл", 
            command=self.save_results
        )
        self.save_button.pack(side=tk.LEFT, padx=5)
        
        # Инициализируем список найденных слов
        self.found_words = []
        
        # Делаем колонки растягиваемыми
        file_frame.columnconfigure(0, weight=1)
        search_frame.columnconfigure(0, weight=1)
    
    def browse_file(self):
        """Открывает диалог выбора файла"""
        filetypes = [("Text files", "*.txt")]
        
        filename = filedialog.askopenfilename(filetypes=filetypes)
        if filename:
            self.file_var.set(filename)
    
    def start_search(self):
        """Запускает поиск слов в отдельном потоке"""
        file_path = self.file_var.get()
        search_word = self.search_var.get().lower()
        
        if not file_path:
            messagebox.showerror("Ошибка", "Выберите файл для поиска!")
            return
        
        if not search_word:
            messagebox.showerror("Ошибка", "Введите слово для поиска!")
            return
        
        # Проверяем существование файла
        if not os.path.exists(file_path):
            messagebox.showerror("Ошибка", "Указанный файл не существует!")
            return
        
        # Очищаем результаты
        self.result_text.delete(1.0, tk.END)
        self.result_text.insert(tk.END, "Выполняется поиск...\n")
        
        # Запускаем поиск в отдельном потоке
        search_thread = threading.Thread(
            target=self.perform_search,
            args=(file_path, search_word)
        )
        search_thread.daemon = True
        search_thread.start()
    
    def perform_search(self, file_path, search_word):
        """Выполняет поиск слов в файле"""
        try:
            # Создаем пустой список для хранения найденных слов
            found_words = []
            
            # Открываем выбранный файл для чтения
            with open(file_path, 'r', encoding='utf-8') as file:
                # Прочитываем файл построчно
                for line in file:
                    # Разделяем строку на слова
                    words = line.split()
                    for word in words:
                        # Очищаем слово от кавычек и других символов для поиска
                        clean_word = word
                        # Удаляем кавычки и другую пунктуацию
                        for char in '"\'",.;:!?…()[]{}':
                            clean_word = clean_word.replace(char, '')
                            
                        # Проверяем, содержит ли слово искомую часть (без учета кавычек)
                        if search_word in clean_word.lower():
                            # Добавляем исходное (неочищенное) слово в список
                            found_words.append(word)
            
            # Убираем дубликаты и сортируем список
            found_words = list(set(found_words))
            found_words.sort()
            
            # Сохраняем найденные слова для возможного сохранения в файл
            self.found_words = found_words
            
            # Форматируем и выводим результаты
            self.after(100, lambda: self.display_search_results(found_words))
            
        except Exception as e:
            # Обрабатываем ошибки
            error_message = f"Ошибка при поиске: {str(e)}"
            self.after(100, lambda: self.result_text.delete(1.0, tk.END))
            self.after(100, lambda: self.result_text.insert(tk.END, error_message))
    
    def display_search_results(self, found_words):
        """Отображает результаты поиска"""
        self.result_text.delete(1.0, tk.END)
        
        if found_words:
            # Форматируем список найденных слов
            formatted_results = self.format_results(found_words)
            self.result_text.insert(tk.END, f'Найденные слова:\n{formatted_results}')
        else:
            self.result_text.insert(tk.END, "Слово не найдено.")
    
    def format_results(self, words):
        """Форматирует список найденных слов"""
        # Форматируем список найденных слов так, чтобы не более 5 слов было в строке, каждое слово в кавычках
        formatted_results = []
        temp = []
        for word in words:
            temp.append(f'"{word}"')
            if len(temp) == 5:
                formatted_results.append(", ".join(temp))
                temp = []
        if temp:
            formatted_results.append(", ".join(temp))
        return "\n".join(formatted_results)
    
    def clear_search(self):
        """Очищает поля поиска и результаты"""
        self.search_var.set("")
        self.result_text.delete(1.0, tk.END)
    
    def copy_results(self):
        """Копирует результаты поиска в буфер обмена"""
        result_text = self.result_text.get(1.0, tk.END)
        self.clipboard_clear()
        self.clipboard_append(result_text)
        messagebox.showinfo("Копирование", "Результаты скопированы в буфер обмена")
        
    def save_results(self):
        """Сохраняет результаты поиска в файл без кавычек"""
        if not self.found_words:
            messagebox.showinfo("Информация", "Нет результатов для сохранения.")
            return
            
        # Открываем диалог сохранения файла
        file_path = filedialog.asksaveasfilename(
            defaultextension=".txt",
            filetypes=[("Текстовые файлы", "*.txt"), ("Все файлы", "*.*")]
        )
        
        if not file_path:
            return
            
        try:
            # Сохраняем найденные слова в файл, по одному слову на строку, без кавычек
            with open(file_path, 'w', encoding='utf-8') as f:
                for word in self.found_words:
                    f.write(word + '\n')
                    
            messagebox.showinfo("Сохранение", f"Результаты успешно сохранены в файл:\n{file_path}")
            
        except Exception as e:
            messagebox.showerror("Ошибка сохранения", f"Не удалось сохранить результаты: {str(e)}")


class ColorAnalysisFrame(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        
        # Создаем элементы интерфейса анализа по списку слов
        self.create_widgets()
    
    def create_widgets(self):
        # Фрейм для выбора файла списка слов
        colors_frame = ttk.LabelFrame(self, text="Список слов")
        colors_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        ttk.Label(colors_frame, text="Файл со списком слов:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        
        self.colors_file_var = tk.StringVar()
        self.colors_file_entry = ttk.Entry(colors_frame, textvariable=self.colors_file_var, width=50)
        self.colors_file_entry.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        
        self.colors_file_button = ttk.Button(
            colors_frame, text="Обзор", 
            command=lambda: self.browse_file(self.colors_file_var)
        )
        self.colors_file_button.grid(row=0, column=2, padx=5, pady=5)
        
        # Фрейм для выбора текстового файла для анализа
        text_frame = ttk.LabelFrame(self, text="Файл для анализа")
        text_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        ttk.Label(text_frame, text="Текстовый файл:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        
        self.text_file_var = tk.StringVar()
        self.text_file_entry = ttk.Entry(text_frame, textvariable=self.text_file_var, width=50)
        self.text_file_entry.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        
        self.text_file_button = ttk.Button(
            text_frame, text="Обзор", 
            command=lambda: self.browse_file(self.text_file_var)
        )
        self.text_file_button.grid(row=0, column=2, padx=5, pady=5)
        
        # Кнопка запуска анализа
        self.analyze_button = ttk.Button(
            self, text="Запустить анализ", 
            command=self.start_analysis
        )
        self.analyze_button.pack(pady=10)
        
        # Индикатор прогресса
        self.progress = ttk.Progressbar(self, orient="horizontal", length=300, mode="indeterminate")
        self.progress.pack(pady=5, fill="x", padx=10)
        
        # Текстовое поле для вывода статуса операции
        self.status_text = tk.Text(self, height=10, width=70, wrap="word")
        self.status_text.pack(padx=10, pady=10, fill="both", expand=True)
        
        # Настройка scrollbar для текстового поля
        scrollbar = ttk.Scrollbar(self.status_text, command=self.status_text.yview)
        scrollbar.pack(side="right", fill="y")
        self.status_text.config(yscrollcommand=scrollbar.set)
        
        # Делаем колонки растягиваемыми
        colors_frame.columnconfigure(1, weight=1)
        text_frame.columnconfigure(1, weight=1)
    
    def browse_file(self, var):
        """Открывает диалог выбора файла и устанавливает значение в переменную"""
        filetypes = [("Text files", "*.txt")]
        
        filename = filedialog.askopenfilename(filetypes=filetypes)
        if filename:
            var.set(filename)
    
    def start_analysis(self):
        """Запускает анализ текста в отдельном потоке"""
        colors_filename = self.colors_file_var.get()
        text_filename = self.text_file_var.get()
        
        if not colors_filename:
            messagebox.showerror("Ошибка", "Выберите файл со списком слов!")
            return
        
        if not text_filename:
            messagebox.showerror("Ошибка", "Выберите файл для анализа!")
            return
        
        # Проверяем существование файлов
        if not os.path.exists(colors_filename):
            messagebox.showerror("Ошибка", "Указанный файл со списком слов не существует!")
            return
        
        if not os.path.exists(text_filename):
            messagebox.showerror("Ошибка", "Указанный файл для анализа не существует!")
            return
        
        # Очищаем поле статуса
        self.status_text.delete(1.0, tk.END)
        self.update_status("Начало анализа...")
        
        # Запускаем прогресс-бар
        self.progress.start()
        
        # Запускаем анализ в отдельном потоке
        analyze_thread = threading.Thread(
            target=self.perform_analysis,
            args=(colors_filename, text_filename)
        )
        analyze_thread.daemon = True
        analyze_thread.start()
    
    def perform_analysis(self, colors_filename, text_filename):
        """Выполняет анализ текста"""
        try:
            # Читаем список слов из файла
            self.update_status("Чтение списка слов...")
            with open(colors_filename, 'r', encoding='utf-8') as f:
                colors = [line.strip() for line in f]
            
            # Читаем текстовый файл
            self.update_status("Чтение текстового файла...")
            with open(text_filename, 'r', encoding='utf-8') as f:
                text = f.read()
            
            # Поиск предложений с указанными словами
            self.update_status("Поиск предложений со словами из списка...")
            sentences = text.split('.')
            sentences_with_colors = []
            
            for sentence in tqdm(sentences, desc='Поиск'):
                for color in colors:
                    # Экранируем специальные символы для безопасного использования в регулярном выражении
                    escaped_color = re.escape(color)
                    try:
                        if re.search(r'\b{}\b'.format(escaped_color), sentence, re.IGNORECASE):
                            sentence_words = sentence.strip().split()
                            formatted_sentence = ""
                            word_count = 0
                            
                            for word in sentence_words:
                                if word_count + len(word.split('\n')) <= 20:
                                    formatted_sentence += word + " "
                                    word_count += len(word.split('\n'))
                                else:
                                    sentences_with_colors.append(formatted_sentence.strip() + '.')
                                    formatted_sentence = word + " "
                                    word_count = len(word.split('\n'))
                                    
                            if formatted_sentence.strip():
                                sentences_with_colors.append(formatted_sentence.strip() + '.')
                            break
                    except Exception as regex_error:
                        self.update_status(f"Ошибка в обработке слова '{color}': {str(regex_error)}")
                        continue
            
            # Extracting file names without extensions
            text_filename_without_ext = os.path.splitext(os.path.basename(text_filename))[0]
            colors_filename_without_ext = os.path.splitext(os.path.basename(colors_filename))[0]
            
            # Constructing the output file name
            output_file_name = f"{text_filename_without_ext}_{colors_filename_without_ext}.txt"
            
            # Сохраняем результаты анализа
            self.update_status("Сохранение результатов анализа...")
            with open(output_file_name, 'w', encoding='utf-8') as f:
                for sentence in tqdm(sentences_with_colors, desc='Запись в файл'):
                    f.write(sentence + '\n')
            
            # Обновляем статус в основном потоке
            self.after(100, lambda: self.analysis_completed(output_file_name))
            
        except Exception as e:
            # Обрабатываем ошибки
            error_message = f"Ошибка при анализе: {str(e)}"
            self.after(100, lambda: self.update_status(error_message))
            self.after(100, lambda: self.progress.stop())
    
    def analysis_completed(self, output_file):
        """Обрабатывает завершение анализа"""
        self.progress.stop()
        self.update_status(f"Анализ завершен успешно!\nРезультат сохранен в: {output_file}")
        
        if messagebox.askyesno("Анализ завершен", 
                              "Анализ завершен успешно. Открыть результаты?"):
            # Открываем файл в блокноте или программе по умолчанию
            os.system(f'notepad.exe "{output_file}"')
    
    def update_status(self, message):
        """Обновляет текстовое поле статуса"""
        self.status_text.insert(tk.END, message + "\n")
        self.status_text.see(tk.END)  # Прокрутка вниз


class TextAnalyzerTab(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        
        # Создаем вложенные вкладки для разных типов анализа
        self.create_widgets()
    
    def create_widgets(self):
        # Создаем систему вложенных вкладок
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill='both', expand=True, padx=10, pady=10)
        
        # Вкладка поиска слов
        self.word_search_frame = WordSearchFrame(self.notebook)
        self.notebook.add(self.word_search_frame, text="Поиск слов")
        
        # Вкладка анализа цветов в тексте
        self.color_analysis_frame = ColorAnalysisFrame(self.notebook)
        self.notebook.add(self.color_analysis_frame, text="Анализ по списку слов")


class TextSplitterTab(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        
        # Создаем элементы интерфейса
        self.create_widgets()
    
    def create_widgets(self):
        # Фрейм для выбора файла
        file_frame = ttk.LabelFrame(self, text="Выбор файла")
        file_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        # Поле для отображения выбранного файла
        self.file_var = tk.StringVar()
        self.file_entry = ttk.Entry(file_frame, textvariable=self.file_var, width=60)
        self.file_entry.grid(row=0, column=0, padx=5, pady=5, sticky="ew")
        
        # Кнопка для выбора файла
        self.browse_button = ttk.Button(file_frame, text="Обзор", command=self.browse_file)
        self.browse_button.grid(row=0, column=1, padx=5, pady=5)
        
        # Фрейм для настройки разделения
        settings_frame = ttk.LabelFrame(self, text="Настройки разделения")
        settings_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        ttk.Label(settings_frame, text="Шаблон для разделения:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        
        # Варианты разделения
        self.split_type = tk.StringVar(value="number")
        
        self.number_radio = ttk.Radiobutton(
            settings_frame, text="По 4-значным числам", 
            variable=self.split_type, value="number"
        )
        self.number_radio.grid(row=1, column=0, padx=5, pady=5, sticky="w")
        
        self.custom_radio = ttk.Radiobutton(
            settings_frame, text="По шаблону:", 
            variable=self.split_type, value="custom"
        )
        self.custom_radio.grid(row=2, column=0, padx=5, pady=5, sticky="w")
        
        self.pattern_var = tk.StringVar()
        self.pattern_entry = ttk.Entry(settings_frame, textvariable=self.pattern_var, width=40)
        self.pattern_entry.grid(row=2, column=1, padx=5, pady=5, sticky="ew")
        
        # Кнопка для запуска разделения текста
        self.split_button = ttk.Button(
            self, text="Разделить текст", 
            command=self.start_splitting
        )
        self.split_button.pack(pady=10)
        
        # Индикатор прогресса
        self.progress = ttk.Progressbar(self, orient="horizontal", length=300, mode="indeterminate")
        self.progress.pack(pady=5, fill="x", padx=10)
        
        # Текстовое поле для вывода статуса операции
        self.status_text = tk.Text(self, height=10, width=70, wrap="word")
        self.status_text.pack(padx=10, pady=10, fill="both", expand=True)
        
        # Настройка scrollbar для текстового поля
        scrollbar = ttk.Scrollbar(self.status_text, command=self.status_text.yview)
        scrollbar.pack(side="right", fill="y")
        self.status_text.config(yscrollcommand=scrollbar.set)
        
        # Делаем первую колонку растягиваемой
        file_frame.columnconfigure(0, weight=1)
        settings_frame.columnconfigure(1, weight=1)
    
    def browse_file(self):
        """Открывает диалог выбора файла"""
        filetypes = [
            ("Текстовые файлы", "*.txt"), 
            ("Все файлы", "*.*")
        ]
        
        filename = filedialog.askopenfilename(filetypes=filetypes)
        if filename:
            self.file_var.set(filename)
    
    def start_splitting(self):
        """Запускает процесс разделения текста в отдельном потоке"""
        file_path = self.file_var.get()
        if not file_path:
            messagebox.showerror("Ошибка", "Выберите файл для разделения!")
            return
        
        # Проверяем существование файла
        if not os.path.exists(file_path):
            messagebox.showerror("Ошибка", "Указанный файл не существует!")
            return
        
        # Проверяем настройки разделения
        split_type = self.split_type.get()
        
        if split_type == "custom" and not self.pattern_var.get():
            messagebox.showerror("Ошибка", "Введите шаблон для разделения!")
            return
        
        # Очищаем поле статуса
        self.status_text.delete(1.0, tk.END)
        self.update_status("Начало разделения текста...")
        
        # Запускаем прогресс-бар
        self.progress.start()
        
        # Запускаем разделение в отдельном потоке
        split_thread = threading.Thread(
            target=self.perform_splitting,
            args=(file_path, split_type, self.pattern_var.get())
        )
        split_thread.daemon = True
        split_thread.start()
    
    def perform_splitting(self, file_path, split_type, custom_pattern):
        """Выполняет разделение текста"""
        try:
            # Определяем шаблон для разделения
            if split_type == "number":
                pattern = r'^\d{4}$'
            else:
                pattern = custom_pattern
            
            self.update_status(f"Использую шаблон: {pattern}")
            
            # Словарь для хранения разделов
            sections = {}
            
            # Чтение и обработка входного файла
            self.update_status("Чтение и обработка файла...")
            
            with open(file_path, 'r', encoding='utf-8') as file:
                current_section = None
                
                for line in file:
                    line = line.strip()
                    
                    if split_type == "number":
                        # Для 4-значных чисел используем простую проверку
                        if line.isdigit() and len(line) == 4:
                            current_section = line
                            sections[current_section] = []
                        elif current_section is not None:
                            sections[current_section].append(line)
                    else:
                        # Для пользовательского шаблона используем регулярные выражения
                        if re.match(pattern, line):
                            current_section = line
                            sections[current_section] = []
                        elif current_section is not None:
                            sections[current_section].append(line)
            
            # Если нет разделов, выводим ошибку
            if not sections:
                self.after(100, lambda: self.update_status("Входной файл не содержит разделов по заданному шаблону!"))
                self.after(100, lambda: self.progress.stop())
                return
            
            # Сохраняем разделы в отдельные файлы
            self.update_status(f"Сохранение {len(sections)} разделов в отдельные файлы...")
            
            output_dir = os.path.dirname(file_path)
            base_name = os.path.splitext(os.path.basename(file_path))[0]
            
            saved_files = []
            
            for section, content in sections.items():
                # Формируем имя выходного файла
                output_file = os.path.join(output_dir, f"{base_name}_{section}.txt")
                
                # Сохраняем раздел в файл
                with open(output_file, 'w', encoding='utf-8') as file:
                    file.write('\n'.join(content))
                
                saved_files.append(output_file)
            
            # Обновляем статус в основном потоке
            self.after(100, lambda: self.splitting_completed(saved_files))
            
        except Exception as e:
            # Обрабатываем ошибки
            error_message = f"Ошибка при разделении текста: {str(e)}"
            self.after(100, lambda: self.update_status(error_message))
            self.after(100, lambda: self.progress.stop())
    
    def splitting_completed(self, saved_files):
        """Обрабатывает завершение разделения текста"""
        self.progress.stop()
        self.update_status(f"Разделение текста завершено успешно!")
        self.update_status(f"Создано файлов: {len(saved_files)}")
        
        for file in saved_files[:5]:  # Показываем только первые 5 файлов
            self.update_status(f"- {os.path.basename(file)}")
        
        if len(saved_files) > 5:
            self.update_status(f"... и еще {len(saved_files) - 5} файлов")
        
        if messagebox.askyesno("Разделение завершено", 
                              "Разделение завершено успешно. Открыть папку с результатами?"):
            # Открываем папку в проводнике
            os.system(f'explorer "{os.path.dirname(saved_files[0])}"')
    
    def update_status(self, message):
        """Обновляет текстовое поле статуса"""
        self.status_text.insert(tk.END, message + "\n")
        self.status_text.see(tk.END)  # Прокрутка вниз


# Добавляем класс статистического анализа текста
class FrequencyAnalysisFrame(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        
        # Создаем элементы интерфейса
        self.create_widgets()
    
    def create_widgets(self):
        # Фрейм для выбора файла
        file_frame = ttk.LabelFrame(self, text="Выбор файла")
        file_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        # Поле для отображения выбранного файла
        self.file_var = tk.StringVar()
        self.file_entry = ttk.Entry(file_frame, textvariable=self.file_var, width=60)
        self.file_entry.grid(row=0, column=0, padx=5, pady=5, sticky="ew")
        
        # Кнопка для выбора файла
        self.browse_button = ttk.Button(file_frame, text="Обзор", command=self.browse_file)
        self.browse_button.grid(row=0, column=1, padx=5, pady=5)
        
        # Фрейм для настроек анализа
        settings_frame = ttk.LabelFrame(self, text="Настройки анализа")
        settings_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        # Настройка минимальной длины слова
        ttk.Label(settings_frame, text="Минимальная длина слова:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        
        self.min_length_var = tk.IntVar(value=3)
        self.min_length_spinbox = ttk.Spinbox(settings_frame, from_=1, to=20, increment=1, textvariable=self.min_length_var, width=5)
        self.min_length_spinbox.grid(row=0, column=1, padx=5, pady=5, sticky="w")
        
        # Настройка количества слов для отображения
        ttk.Label(settings_frame, text="Количество слов в топе:").grid(row=1, column=0, padx=5, pady=5, sticky="w")
        
        self.top_words_var = tk.IntVar(value=20)
        self.top_words_spinbox = ttk.Spinbox(settings_frame, from_=5, to=100, increment=5, textvariable=self.top_words_var, width=5)
        self.top_words_spinbox.grid(row=1, column=1, padx=5, pady=5, sticky="w")
        
        # Опция исключения стоп-слов
        self.exclude_stopwords_var = tk.BooleanVar(value=True)
        self.exclude_stopwords_check = ttk.Checkbutton(
            settings_frame, text="Исключить стоп-слова", 
            variable=self.exclude_stopwords_var
        )
        self.exclude_stopwords_check.grid(row=2, column=0, columnspan=2, padx=5, pady=5, sticky="w")
        
        # Кнопка для анализа
        self.analyze_button = ttk.Button(
            self, text="Выполнить анализ", 
            command=self.start_analysis
        )
        self.analyze_button.pack(pady=10)
        
        # Индикатор прогресса
        self.progress = ttk.Progressbar(self, orient="horizontal", length=300, mode="indeterminate")
        self.progress.pack(pady=5, fill="x", padx=10)
        
        # Фрейм для результатов
        results_frame = ttk.LabelFrame(self, text="Результаты анализа")
        results_frame.pack(fill="both", padx=10, pady=10, expand=True)
        
        # Создаем фрейм для графика
        self.plot_frame = ttk.Frame(results_frame)
        self.plot_frame.pack(side=tk.LEFT, fill="both", expand=True, padx=5, pady=5)
        
        # Создаем фрейм для таблицы
        self.table_frame = ttk.Frame(results_frame)
        self.table_frame.pack(side=tk.RIGHT, fill="both", expand=False, padx=5, pady=5)
        
        # Создаем таблицу результатов
        columns = ("слово", "частота", "процент")
        self.results_tree = ttk.Treeview(self.table_frame, columns=columns, show="headings", height=15)
        
        # Настройка заголовков
        self.results_tree.heading("слово", text="Слово")
        self.results_tree.heading("частота", text="Частота")
        self.results_tree.heading("процент", text="Процент")
        
        # Настройка столбцов
        self.results_tree.column("слово", width=120)
        self.results_tree.column("частота", width=80, anchor="center")
        self.results_tree.column("процент", width=80, anchor="center")
        
        # Добавляем scrollbar для таблицы
        tree_scrollbar = ttk.Scrollbar(self.table_frame, orient="vertical", command=self.results_tree.yview)
        self.results_tree.configure(yscrollcommand=tree_scrollbar.set)
        
        # Размещаем элементы
        self.results_tree.pack(side=tk.LEFT, fill="both", expand=True)
        tree_scrollbar.pack(side=tk.RIGHT, fill="y")
        
        # Кнопка для сохранения результатов
        self.save_button = ttk.Button(
            results_frame, text="Сохранить результаты", 
            command=self.save_results,
            state=tk.DISABLED
        )
        self.save_button.pack(pady=5)
        
        # Делаем первую колонку растягиваемой
        file_frame.columnconfigure(0, weight=1)
        
        # Русские стоп-слова
        self.russian_stopwords = [
            "а", "без", "более", "бы", "был", "была", "были", "было", "быть", "в", 
            "вам", "вас", "весь", "во", "вот", "все", "всего", "всех", "вы", "где", 
            "да", "даже", "для", "до", "его", "ее", "если", "есть", "еще", "же", 
            "за", "здесь", "и", "из", "или", "им", "их", "к", "как", "ко", 
            "когда", "кто", "ли", "либо", "мне", "может", "мы", "на", "надо", "наш", 
            "не", "него", "нее", "нет", "ни", "но", "ну", "о", "об", "однако", 
            "он", "она", "они", "оно", "от", "очень", "по", "под", "при", "с", 
            "со", "так", "также", "такой", "там", "те", "тем", "то", "того", "тоже", 
            "той", "только", "том", "ты", "у", "уже", "хотя", "чего", "чей", "чем", 
            "что", "чтобы", "этого", "этой", "этом", "этот", "эту", "я"
        ]
        
        # Данные для графика и таблицы
        self.analysis_results = None
    
    def browse_file(self):
        """Открывает диалог выбора файла"""
        filetypes = [
            ("Текстовые файлы", "*.txt"), 
            ("Все файлы", "*.*")
        ]
        
        filename = filedialog.askopenfilename(filetypes=filetypes)
        if filename:
            self.file_var.set(filename)
    
    def start_analysis(self):
        """Запускает частотный анализ в отдельном потоке"""
        file_path = self.file_var.get()
        if not file_path:
            messagebox.showerror("Ошибка", "Выберите файл для анализа!")
            return
        
        # Проверяем существование файла
        if not os.path.exists(file_path):
            messagebox.showerror("Ошибка", "Указанный файл не существует!")
            return
        
        # Получаем настройки
        try:
            min_length = int(self.min_length_var.get())
            top_words = int(self.top_words_var.get())
            
            if min_length <= 0:
                raise ValueError("Минимальная длина слова должна быть положительным числом")
                
            if top_words <= 0:
                raise ValueError("Количество слов в топе должно быть положительным числом")
                
        except ValueError as e:
            messagebox.showerror("Ошибка", f"Некорректные настройки анализа: {str(e)}")
            return
        
        # Очищаем предыдущие результаты
        self.clear_results()
        
        # Запускаем прогресс-бар
        self.progress.start()
        
        # Блокируем кнопку анализа
        self.analyze_button.config(state=tk.DISABLED)
        
        # Запускаем анализ в отдельном потоке
        analyze_thread = threading.Thread(
            target=self.perform_analysis,
            args=(file_path, min_length, top_words, self.exclude_stopwords_var.get())
        )
        analyze_thread.daemon = True
        analyze_thread.start()
    
    def perform_analysis(self, file_path, min_length, top_words, exclude_stopwords):
        """Выполняет частотный анализ текста"""
        try:
            # Читаем файл
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
            
            # Приводим текст к нижнему регистру
            text = text.lower()
            
            # Удаляем пунктуацию и заменяем её пробелами
            text = re.sub(r'[^\w\s]', ' ', text)
            
            # Разбиваем на слова
            words = text.split()
            
            # Фильтруем слова по длине
            words = [word for word in words if len(word) >= min_length]
            
            # Исключаем стоп-слова если нужно
            if exclude_stopwords:
                words = [word for word in words if word not in self.russian_stopwords]
            
            # Подсчитываем частоту слов
            word_counts = Counter(words)
            
            # Получаем общее количество слов
            total_words = len(words)
            
            # Получаем наиболее часто встречающиеся слова
            most_common_words = word_counts.most_common(top_words)
            
            # Создаем список для хранения результатов анализа: (слово, частота, процент)
            results = []
            for word, count in most_common_words:
                percent = (count / total_words) * 100
                results.append((word, count, percent))
            
            # Сохраняем результаты для использования в других методах
            self.analysis_results = results
            
            # Обновляем интерфейс в основном потоке
            self.after(100, lambda: self.update_results(results, total_words))
            
        except Exception as e:
            # Обрабатываем ошибки
            error_message = f"Ошибка при анализе: {str(e)}"
            self.after(100, lambda: messagebox.showerror("Ошибка", error_message))
            self.after(100, lambda: self.progress.stop())
            self.after(100, lambda: self.analyze_button.config(state=tk.NORMAL))
    
    def update_results(self, results, total_words):
        """Обновляет отображение результатов анализа"""
        # Останавливаем прогресс-бар
        self.progress.stop()
        
        # Разблокируем кнопку анализа
        self.analyze_button.config(state=tk.NORMAL)
        
        # Добавляем данные в таблицу
        for word, count, percent in results:
            self.results_tree.insert("", "end", values=(word, count, f"{percent:.2f}%"))
        
        # Строим гистограмму
        self.create_frequency_plot(results)
        
        # Активируем кнопку сохранения
        self.save_button.config(state=tk.NORMAL)
        
        # Вывод общей информации
        messagebox.showinfo("Анализ завершен", 
                          f"Анализ успешно завершен.\nВсего слов в тексте: {total_words}\n"
                          f"Отображено топ {len(results)} слов.")
    
    def create_frequency_plot(self, results):
        """Создает гистограмму частотности слов с адаптивным дизайном"""
        # Очищаем предыдущий график
        for widget in self.plot_frame.winfo_children():
            widget.destroy()
        
        # Подготавливаем данные для графика (берем только 10 самых частых слов для наглядности)
        plot_data = results[:10]
        words = [item[0] for item in plot_data]
        frequencies = [item[1] for item in plot_data]
        
        # Создаем фрейм для графика с поддержкой изменения размера
        plot_container = ttk.Frame(self.plot_frame)
        plot_container.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        # Создаем фигуру и оси с поддержкой адаптивного размера
        fig = Figure(dpi=100)
        ax = fig.add_subplot(111)
        
        # Строим горизонтальную гистограмму
        bars = ax.barh(words, frequencies, color='skyblue')
        
        # Добавляем значения на гистограмму
        for i, bar in enumerate(bars):
            ax.text(bar.get_width() + 1, bar.get_y() + bar.get_height()/2, 
                    f'{frequencies[i]}', va='center')
        
        # Настраиваем график
        ax.set_title('Топ 10 самых частых слов', fontsize=11)
        ax.set_xlabel('Частота', fontsize=10)
        ax.set_ylabel('Слово', fontsize=10)
        
        # Инвертируем ось Y, чтобы самые частые слова были сверху
        ax.invert_yaxis()
        
        # Настраиваем поля для лучшей адаптивности
        fig.tight_layout()
        
        # Создаем холст для отображения графика в Tkinter
        canvas = FigureCanvasTkAgg(fig, master=plot_container)
        canvas_widget = canvas.get_tk_widget()
        canvas_widget.pack(fill=tk.BOTH, expand=True)
        
        # Добавляем панель инструментов для графика с возможностью масштабирования
        from matplotlib.backends.backend_tkagg import NavigationToolbar2Tk
        toolbar_frame = ttk.Frame(plot_container)
        toolbar_frame.pack(fill=tk.X, side=tk.BOTTOM)
        toolbar = NavigationToolbar2Tk(canvas, toolbar_frame)
        toolbar.update()
        
        # Устанавливаем обработчик изменения размера окна
        def on_resize(event):
            # Обновляем размер фигуры при изменении размера контейнера
            fig.tight_layout()
            canvas.draw()
            
        # Привязываем обработчик к событию изменения размера
        plot_container.bind("<Configure>", on_resize)
        
        # Рисуем график
        canvas.draw()
    
    def clear_results(self):
        """Очищает предыдущие результаты анализа"""
        # Очищаем таблицу
        for item in self.results_tree.get_children():
            self.results_tree.delete(item)
        
        # Очищаем график
        for widget in self.plot_frame.winfo_children():
            widget.destroy()
        
        # Сбрасываем результаты
        self.analysis_results = None
        
        # Деактивируем кнопку сохранения
        self.save_button.config(state=tk.DISABLED)
    
    def save_results(self):
        """Сохраняет результаты анализа в файл"""
        if not self.analysis_results:
            return
        
        # Открываем диалог сохранения файла
        file_path = filedialog.asksaveasfilename(
            defaultextension=".txt",
            filetypes=[("Текстовые файлы", "*.txt"), ("Все файлы", "*.*")]
        )
        
        if not file_path:
            return
        
        try:
            with open(file_path, 'w', encoding='utf-8') as file:
                # Записываем заголовок
                file.write("Слово\tЧастота\tПроцент\n")
                
                # Записываем данные
                for word, count, percent in self.analysis_results:
                    file.write(f"{word}\t{count}\t{percent:.2f}%\n")
            
            messagebox.showinfo("Сохранение результатов", 
                              f"Результаты успешно сохранены в файл:\n{file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка сохранения", 
                               f"Не удалось сохранить результаты: {str(e)}")


class WordCloudFrame(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        
        # Создаем элементы интерфейса
        self.create_widgets()
    
    def create_widgets(self):
        # Фрейм для выбора файла
        file_frame = ttk.LabelFrame(self, text="Выбор файла")
        file_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        # Поле для отображения выбранного файла
        self.file_var = tk.StringVar()
        self.file_entry = ttk.Entry(file_frame, textvariable=self.file_var, width=60)
        self.file_entry.grid(row=0, column=0, padx=5, pady=5, sticky="ew")
        
        # Кнопка для выбора файла
        self.browse_button = ttk.Button(file_frame, text="Обзор", command=self.browse_file)
        self.browse_button.grid(row=0, column=1, padx=5, pady=5)
        
        # Создаем frame для настроек с поддержкой прокрутки для адаптивности
        settings_outer_frame = ttk.Frame(self)
        settings_outer_frame.pack(fill="x", padx=10, pady=10, expand=False)
        
        # Создаем Canvas и scrollbar для прокрутки
        settings_canvas = tk.Canvas(settings_outer_frame, highlightthickness=0)
        scrollbar = ttk.Scrollbar(settings_outer_frame, orient="vertical", command=settings_canvas.yview)
        
        # Настраиваем Canvas для прокрутки
        settings_canvas.configure(yscrollcommand=scrollbar.set)
        settings_canvas.pack(side=tk.LEFT, fill="both", expand=True)
        scrollbar.pack(side=tk.RIGHT, fill="y")
        
        # Фрейм для настроек внутри Canvas
        settings_frame = ttk.LabelFrame(settings_canvas, text="Настройки")
        settings_window = settings_canvas.create_window((0, 0), window=settings_frame, anchor="nw", tags="settings_frame")
        
        # Настройка минимальной длины слова
        ttk.Label(settings_frame, text="Минимальная длина слова:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        
        self.min_length_var = tk.IntVar(value=3)
        self.min_length_spinbox = ttk.Spinbox(settings_frame, from_=1, to=20, increment=1, textvariable=self.min_length_var, width=5)
        self.min_length_spinbox.grid(row=0, column=1, padx=5, pady=5, sticky="w")
        
        # Настройка максимального количества слов
        ttk.Label(settings_frame, text="Максимальное количество слов:").grid(row=1, column=0, padx=5, pady=5, sticky="w")
        
        self.max_words_var = tk.IntVar(value=200)
        self.max_words_spinbox = ttk.Spinbox(settings_frame, from_=50, to=1000, increment=50, textvariable=self.max_words_var, width=5)
        self.max_words_spinbox.grid(row=1, column=1, padx=5, pady=5, sticky="w")
        
        # Выбор цветовой схемы
        ttk.Label(settings_frame, text="Цветовая схема:").grid(row=2, column=0, padx=5, pady=5, sticky="w")
        
        self.colormap_var = tk.StringVar(value="viridis")
        self.colormap_combo = ttk.Combobox(settings_frame, textvariable=self.colormap_var, width=15)
        self.colormap_combo['values'] = ('viridis', 'plasma', 'inferno', 'magma', 'cividis', 
                                         'Blues', 'Greens', 'Reds', 'YlOrBr', 'RdPu',
                                         'rainbow', 'hsv', 'jet')
        self.colormap_combo.grid(row=2, column=1, padx=5, pady=5, sticky="w")
        
        # Настройка размеров облака
        ttk.Label(settings_frame, text="Ширина облака:").grid(row=3, column=0, padx=5, pady=5, sticky="w")
        self.width_var = tk.IntVar(value=800)
        self.width_spinbox = ttk.Spinbox(settings_frame, from_=400, to=2000, increment=100, textvariable=self.width_var, width=5)
        self.width_spinbox.grid(row=3, column=1, padx=5, pady=5, sticky="w")
        
        ttk.Label(settings_frame, text="Высота облака:").grid(row=4, column=0, padx=5, pady=5, sticky="w")
        self.height_var = tk.IntVar(value=500)
        self.height_spinbox = ttk.Spinbox(settings_frame, from_=300, to=1500, increment=100, textvariable=self.height_var, width=5)
        self.height_spinbox.grid(row=4, column=1, padx=5, pady=5, sticky="w")
        
        # Опция исключения стоп-слов
        self.exclude_stopwords_var = tk.BooleanVar(value=True)
        self.exclude_stopwords_check = ttk.Checkbutton(
            settings_frame, text="Исключить стоп-слова", 
            variable=self.exclude_stopwords_var
        )
        self.exclude_stopwords_check.grid(row=5, column=0, columnspan=2, padx=5, pady=5, sticky="w")
        
        # Настройка масштабирования Canvas при изменении размера фрейма настроек
        def on_frame_configure(event):
            settings_canvas.configure(scrollregion=settings_canvas.bbox("all"))
            settings_canvas.itemconfig(settings_window, width=event.width)
        
        settings_frame.bind("<Configure>", on_frame_configure)
        settings_frame.columnconfigure(0, weight=1)
        
        # Кнопка для генерации облака слов
        self.generate_button = ttk.Button(
            self, text="Создать облако слов", 
            command=self.start_generation
        )
        self.generate_button.pack(pady=10)
        
        # Индикатор прогресса
        self.progress = ttk.Progressbar(self, orient="horizontal", length=300, mode="indeterminate")
        self.progress.pack(pady=5, fill="x", padx=10)
        
        # Фрейм для отображения облака слов
        self.wordcloud_frame = ttk.LabelFrame(self, text="Облако слов")
        self.wordcloud_frame.pack(fill="both", padx=10, pady=10, expand=True)
        
        # Фрейм для размещения изображения
        self.image_frame = ttk.Frame(self.wordcloud_frame)
        self.image_frame.pack(fill="both", expand=True, padx=5, pady=5)
        
        # Кнопка для сохранения облака слов
        self.save_button = ttk.Button(
            self.wordcloud_frame, text="Сохранить изображение", 
            command=self.save_wordcloud,
            state=tk.DISABLED
        )
        self.save_button.pack(pady=5)
        
        # Делаем первую колонку растягиваемой
        file_frame.columnconfigure(0, weight=1)
        
        # Русские стоп-слова
        self.russian_stopwords = [
            "а", "без", "более", "бы", "был", "была", "были", "было", "быть", "в", 
            "вам", "вас", "весь", "во", "вот", "все", "всего", "всех", "вы", "где", 
            "да", "даже", "для", "до", "его", "ее", "если", "есть", "еще", "же", 
            "за", "здесь", "и", "из", "или", "им", "их", "к", "как", "ко", 
            "когда", "кто", "ли", "либо", "мне", "может", "мы", "на", "надо", "наш", 
            "не", "него", "нее", "нет", "ни", "но", "ну", "о", "об", "однако", 
            "он", "она", "они", "оно", "от", "очень", "по", "под", "при", "с", 
            "со", "так", "также", "такой", "там", "те", "тем", "то", "того", "тоже", 
            "той", "только", "том", "ты", "у", "уже", "хотя", "чего", "чей", "чем", 
            "что", "чтобы", "этого", "этой", "этом", "этот", "эту", "я"
        ]
        
        # Сохраняем последнее созданное облако слов
        self.wordcloud_image = None
    
    def browse_file(self):
        """Открывает диалог выбора файла"""
        filetypes = [
            ("Текстовые файлы", "*.txt"), 
            ("Все файлы", "*.*")
        ]
        
        filename = filedialog.askopenfilename(filetypes=filetypes)
        if filename:
            self.file_var.set(filename)
    
    def start_generation(self):
        """Запускает генерацию облака слов в отдельном потоке"""
        file_path = self.file_var.get()
        if not file_path:
            messagebox.showerror("Ошибка", "Выберите файл для анализа!")
            return
        
        # Проверяем существование файла
        if not os.path.exists(file_path):
            messagebox.showerror("Ошибка", "Указанный файл не существует!")
            return
        
        # Получаем настройки
        try:
            min_length = int(self.min_length_var.get())
            max_words = int(self.max_words_var.get())
            colormap = self.colormap_var.get()
            
            if min_length <= 0:
                raise ValueError("Минимальная длина слова должна быть положительным числом")
                
            if max_words <= 0:
                raise ValueError("Максимальное количество слов должно быть положительным числом")
                
        except ValueError as e:
            messagebox.showerror("Ошибка", f"Некорректные настройки: {str(e)}")
            return
        
        # Очищаем предыдущие результаты
        self.clear_results()
        
        # Запускаем прогресс-бар
        self.progress.start()
        
        # Блокируем кнопку генерации
        self.generate_button.config(state=tk.DISABLED)
        
        # Запускаем генерацию в отдельном потоке
        generate_thread = threading.Thread(
            target=self.generate_wordcloud,
            args=(file_path, min_length, max_words, colormap, self.exclude_stopwords_var.get())
        )
        generate_thread.daemon = True
        generate_thread.start()
    
    def generate_wordcloud(self, file_path, min_length, max_words, colormap, exclude_stopwords):
        """Генерирует облако слов из текста"""
        try:
            # Читаем файл
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
            
            # Приводим текст к нижнему регистру
            text = text.lower()
            
            # Определяем стоп-слова
            stopwords = self.russian_stopwords if exclude_stopwords else None
            
            # Получаем заданные пользователем размеры
            width = self.width_var.get()
            height = self.height_var.get()
            
            # Генерируем облако слов с настраиваемыми параметрами
            wordcloud = WordCloud(
                width=width, 
                height=height, 
                max_words=max_words,
                background_color='white',
                colormap=colormap,
                min_word_length=min_length,
                stopwords=stopwords,
                collocations=False,
                # Указываем шрифт, который поддерживает кириллицу
                font_path=None,  # WordCloud автоматически выберет подходящий шрифт
                prefer_horizontal=0.9,  # Больше горизонтальных слов для лучшей читаемости
                scale=1.5  # Масштабирование для лучшей детализации
            ).generate(text)
            
            # Сохраняем созданное облако слов
            self.wordcloud_image = wordcloud
            
            # Обновляем интерфейс в основном потоке
            self.after(100, lambda: self.display_wordcloud(wordcloud))
            
        except Exception as e:
            # Обрабатываем ошибки
            error_message = f"Ошибка при создании облака слов: {str(e)}"
            self.after(100, lambda: messagebox.showerror("Ошибка", error_message))
            self.after(100, lambda: self.progress.stop())
            self.after(100, lambda: self.generate_button.config(state=tk.NORMAL))
    
    def display_wordcloud(self, wordcloud):
        """Отображает сгенерированное облако слов с адаптивным дизайном"""
        # Останавливаем прогресс-бар
        self.progress.stop()
        
        # Разблокируем кнопку генерации
        self.generate_button.config(state=tk.NORMAL)
        
        # Очищаем предыдущее изображение
        for widget in self.image_frame.winfo_children():
            widget.destroy()
        
        # Создаем адаптивный контейнер для графика
        cloud_container = ttk.Frame(self.image_frame)
        cloud_container.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        # Создаем matplotlib фигуру с адаптивным размером
        fig = Figure(dpi=100)
        ax = fig.add_subplot(111)
        
        # Отображаем облако слов
        ax.imshow(wordcloud, interpolation='bilinear')
        ax.axis('off')  # Скрываем оси
        
        # Настраиваем поля для лучшей адаптивности
        fig.tight_layout()
        
        # Создаем холст для отображения в Tkinter
        canvas = FigureCanvasTkAgg(fig, master=cloud_container)
        canvas_widget = canvas.get_tk_widget()
        canvas_widget.pack(fill=tk.BOTH, expand=True)
        
        # Добавляем панель инструментов для интерактивности
        from matplotlib.backends.backend_tkagg import NavigationToolbar2Tk
        toolbar_frame = ttk.Frame(cloud_container)
        toolbar_frame.pack(fill=tk.X, side=tk.BOTTOM)
        toolbar = NavigationToolbar2Tk(canvas, toolbar_frame)
        toolbar.update()
        
        # Устанавливаем обработчик изменения размера окна
        def on_resize(event):
            # Обновляем размер фигуры при изменении размера контейнера
            fig.tight_layout()
            canvas.draw()
            
        # Привязываем обработчик к событию изменения размера
        cloud_container.bind("<Configure>", on_resize)
        
        # Рисуем облако слов
        canvas.draw()
        
        # Активируем кнопку сохранения
        self.save_button.config(state=tk.NORMAL)
        
        # Показываем сообщение об успешном завершении
        messagebox.showinfo("Облако слов создано", 
                           "Облако слов успешно создано. Используйте панель инструментов для масштабирования и сохранения.")
    
    def clear_results(self):
        """Очищает предыдущие результаты"""
        # Очищаем фрейм с изображением
        for widget in self.image_frame.winfo_children():
            widget.destroy()
        
        # Сбрасываем облако слов
        self.wordcloud_image = None
        
        # Деактивируем кнопку сохранения
        self.save_button.config(state=tk.DISABLED)
    
    def save_wordcloud(self):
        """Сохраняет сгенерированное облако слов в файл"""
        if not self.wordcloud_image:
            return
        
        # Открываем диалог сохранения файла
        file_path = filedialog.asksaveasfilename(
            defaultextension=".png",
            filetypes=[("PNG файлы", "*.png"), ("Все файлы", "*.*")]
        )
        
        if not file_path:
            return
        
        try:
            # Сохраняем изображение
            self.wordcloud_image.to_file(file_path)
            
            messagebox.showinfo("Сохранение облака слов", 
                              f"Облако слов успешно сохранено в файл:\n{file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка сохранения", 
                               f"Не удалось сохранить облако слов: {str(e)}")


class TextStatisticsTab(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        
        # Создаем элементы интерфейса
        self.create_widgets()
    
    def create_widgets(self):
        # Создаем систему вложенных вкладок для разных типов статистики
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill='both', expand=True, padx=10, pady=10)
        
        # Вкладка частотного анализа
        self.frequency_frame = FrequencyAnalysisFrame(self.notebook)
        self.notebook.add(self.frequency_frame, text="Частотный анализ")
        
        # Вкладка облака слов
        self.wordcloud_frame = WordCloudFrame(self.notebook)
        self.notebook.add(self.wordcloud_frame, text="Облако слов")


if __name__ == "__main__":
    app = MultiTextApp()
    app.mainloop()